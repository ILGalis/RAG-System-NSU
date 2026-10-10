
from pathlib import Path

from docx import Document

from app.rag.parsers.base import ParsedDocument


class DOCXParser:

    def parse(self, file_path: str | Path) -> ParsedDocument:
        path = Path(file_path)

        if not path.is_file():
            raise FileNotFoundError(f"DOCX file not found: {path}")

        if path.suffix.lower() != ".docx":
            raise ValueError(
                f"Expected a DOCX file, got: {path.name}"
            )

        try:
            doc = Document(path)
        except Exception as exc:
            raise ValueError(
                f"Could not read DOCX file: {path.name}"
            ) from exc

        parts: list[str] = []

        from docx.table import Table
        from docx.text.paragraph import Paragraph

        for element in doc.iter_inner_content():
            if isinstance(element, Paragraph):
                text = element.text.strip()
                if text:
                    parts.append(text)

            elif isinstance(element, Table):
                for row in element.rows:
                    cells = [
                        cell.text.strip().replace("\n", " / ")
                        for cell in row.cells
                    ]
                    parts.append(" | ".join(cells))

        return ParsedDocument(
            text="\n\n".join(parts),
            source_name=path.name,
            file_type="docx",
            metadata={
                "paragraph_count": str(len(doc.paragraphs)),
                "table_count": str(len(doc.tables)),
            },
        )
